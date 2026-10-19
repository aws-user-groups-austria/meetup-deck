"""A small template with the marker slides of the group's deck, built per test.

The real deck holds organizer emails and photos and is never committed, so the
tests build their own: the same texts the build looks for, the same layout of
a full-slide header picture, a logo on the shoutout, off-slide #Todo boxes,
and one hidden backup slide that must come out untouched.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

EVENT = {
    "version": 1,
    "id": "2026-10-20-aws-meetup-vienna-53",
    "source": {"meetupUrl": "https://www.meetup.com/g/events/1/", "fetchedAt": "2026-10-16T12:00:00.000Z"},
    "title": "AWS Meetup Vienna #53 - hosted by RBI",
    "group": {"id": "aws-ug-vienna", "name": "AWS User Group Vienna", "url": "https://www.meetup.com/g/"},
    "timezone": "Europe/Vienna",
    "doorsOpen": "2026-10-20T17:00:00+02:00",
    "end": "2026-10-20T22:00:00+02:00",
    "venue": {"name": "Raiffeisen Bank International", "address": "Mooslackengasse 12, 1190 Wien, Austria"},
    "sponsors": [{"name": "Raiffeisen Bank International AG", "short": "RBI"}],
    "segments": [
        {"id": "doors", "kind": "networking", "start": "2026-10-20T17:00:00+02:00", "end": "2026-10-20T18:00:00+02:00", "title": "Entry & Networking"},
        {"id": "welcome", "kind": "welcome", "start": "2026-10-20T18:00:00+02:00", "end": "2026-10-20T18:05:00+02:00", "title": "Welcome"},
        {"id": "talk-1", "kind": "talk", "start": "2026-10-20T18:10:00+02:00", "end": "2026-10-20T18:50:00+02:00", "title": "Talk one",
         "talk": {"title": "Talk one", "speakers": [{"name": "Roman Jokel"}, {"name": "Mladen Trampic"}]}},
        {"id": "break-1", "kind": "break", "start": "2026-10-20T18:50:00+02:00", "end": "2026-10-20T19:20:00+02:00", "title": "Networking break"},
        {"id": "talk-2", "kind": "talk", "start": "2026-10-20T19:20:00+02:00", "end": "2026-10-20T20:00:00+02:00", "title": "Talk two",
         "talk": {"title": "Talk two", "speakers": [{"name": "Darko Mesaros"}]}},
        {"id": "closing", "kind": "closing", "start": "2026-10-20T20:00:00+02:00", "end": "2026-10-20T22:00:00+02:00", "title": "Networking"},
    ],
    "links": {"meetup": "https://www.meetup.com/g/events/1/", "feedback": "https://www.meetup.com/g/events/1/", "nextEvents": "https://awscommunity.at/"},
    "notices": [],
}

STILLS = ["Agenda", "SpeakerHandover--talk-1", "Break--break-1", "SpeakerHandover--talk-2", "Closing"]


def png(path: Path, color: str, size: tuple[int, int] = (160, 90)) -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def text(slide, s: str, x: float, y: float, w: float = 9.3, h: float = 0.8, size: int = 26):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.text = s
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(size)
    run.font.name = "PT Sans"
    return box


def todo(slide, s: str):
    return text(slide, f"#Todo:\n{s}", 10.34, 0, 1.8, 1.6, 12)


def make_template(path: Path, tmp: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)  # welcome
    text(s, "Welcome", 0.3, 0.9)
    text(s, "to the AWS User Group Vienna", 0.3, 2.2)
    text(s, "Meetup #52", 3.8, 2.8)
    todo(s, "Change the Number")

    s = prs.slides.add_slide(blank)  # shoutout
    text(s, "Shoutout", 0.3, 0.9)
    text(s, "to Our Food & Venue Sponsors", 0.3, 2.2)
    s.shapes.add_picture(str(png(tmp / "old-logo.png", "pink")), Inches(4.6), Inches(3), Inches(2), Inches(1.5))
    todo(s, "Replace the sponsor logos")

    s = prs.slides.add_slide(blank)  # header
    text(s, "#X Meetup", 0.3, 0.8)
    text(s, "hosted by Y", 0.3, 3.1)
    todo(s, "Update the title slide with the header image")
    s.shapes.add_picture(str(png(tmp / "old-header.png", "gray")), 0, 0, prs.slide_width, prs.slide_height)

    s = prs.slides.add_slide(blank)  # members
    text(s, "Introducing Our User Group Members", 0.3, 0.6)

    s = prs.slides.add_slide(blank)  # hidden backup, must stay as it is
    text(s, "Meetup #47", 3.8, 2.8)
    s._element.set("show", "0")

    s = prs.slides.add_slide(blank)  # see you soon
    text(s, "See you soon", 0.3, 0.9)
    text(s, "at our Meetup #53\x0bOctober 20th", 0.3, 2.2)
    todo(s, "Change the Number")

    prs.save(str(path))
    return path


@pytest.fixture
def setup(tmp_path: Path):
    """(template, event.json, stills dir, header image) in a temp folder."""
    template = make_template(tmp_path / "template.pptx", tmp_path)
    event_dir = tmp_path / "events" / EVENT["id"]
    event_dir.mkdir(parents=True)
    event_file = event_dir / "event.json"
    event_file.write_text(json.dumps(EVENT))
    stills = tmp_path / "stills"
    stills.mkdir()
    for name in STILLS:
        png(stills / f"{name}.png", "purple", (192, 108))
    header = png(tmp_path / "header.jpeg.png", "blue", (160, 90))
    return template, event_file, stills, header
