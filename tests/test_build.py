import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from meetup_deck.build import Inputs, build, evening_slides
from meetup_deck.cli import main


def texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def visible(prs):
    return [s for s in prs.slides if s._element.get("show") != "0"]


def notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def run(setup, tmp_path: Path, header: bool = True):
    template, event_file, stills, header_file = setup
    out = tmp_path / "out.pptx"
    report = build(Inputs(template, event_file, stills, header_file if header else None), out)
    return Presentation(str(out)), report


def test_running_order(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    order = [texts(s)[0] if texts(s) else notes(s).split("  ")[-1].split("\n")[0][:20] for s in visible(prs)]
    assert order[0] == "Welcome"
    assert order[1] == "Thank you"
    assert order[2] == "Meetup #53"
    assert notes(visible(prs)[3]).startswith("Agenda, from https://www.meetup.com/g/events/1/")
    assert order[4] == "Introducing Our User Group Members"
    assert notes(visible(prs)[5]).split("\n")[0] == "18:10-18:50  Talk one"
    # the host's own slot, right after the first talk's slide
    assert order[6] == "Shoutout"
    assert notes(visible(prs)[6]).startswith("Sponsor intro: RBI introduces themselves")
    assert [notes(s).split("\n")[0] for s in visible(prs)[7:10]] == [
        "18:50-19:20  Networking break. Back at 19:20. Next: Talk two at 19:20.",
        "19:20-20:00  Talk two",
        "20:00  Thank the speakers and the sponsor, ask for feedback (https://www.meetup.com/g/events/1/), networking until 22:00.",
    ]
    assert order[10] == "See you soon"
    assert len(visible(prs)) == 11


def test_host_thanked_at_start_and_end_and_its_slot_after_the_first_talk(setup, tmp_path):
    template, event_file, stills, header = setup
    # the real template has a second shoutout at the end
    prs = Presentation(str(template))
    from conftest import png, text
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, "Shoutout", 0.3, 0.9)
    text(s, "to Our Food & Venue Sponsors", 0.3, 2.2)
    s.shapes.add_picture(str(png(tmp_path / "logo2.png", "blue")), 0, 0)
    prs.save(str(template))
    out = tmp_path / "out.pptx"
    build(Inputs(template, event_file, stills, header), out)
    deck = Presentation(str(out))
    slides = visible(deck)
    firsts = [texts(sl)[0] if texts(sl) else "" for sl in slides]
    assert firsts.count("Shoutout") == 1
    hosts = [sl for sl in slides if any(t == "to our Location & Food Sponsor" for t in texts(sl))]
    assert len(hosts) == 2 and all(texts(sl)[0] == "Thank you" for sl in hosts)
    opening, closing = hosts
    assert slides.index(opening) == 1 and slides.index(closing) == len(slides) - 1
    assert notes(opening).startswith("A quick thank you to RBI for hosting.")
    assert texts(closing)[1] == "to our Location & Food Sponsor"
    assert not any("Before the first talk" in t for sl in slides for t in texts(sl))
    slot = slides[firsts.index("Shoutout")]
    assert notes(slides[firsts.index("Shoutout") - 1]).startswith("18:10-18:50  Talk one")
    # the slot shows the sponsor like the template's shoutout: its logo, or its name until there is one
    assert any(sh.shape_type == 13 for sh in slot.shapes) or any("Raiffeisen" in t for t in texts(slot))
    assert texts(slot)[1] == "to Our Food & Venue Sponsors"


def test_welcome_number_and_todos(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    welcome = visible(prs)[0]
    assert "Meetup #53" in texts(welcome)
    assert not any("#Todo" in t for t in texts(welcome))
    run_font = next(sh for sh in welcome.shapes if sh.has_text_frame and sh.text_frame.text == "Meetup #53")
    assert run_font.text_frame.paragraphs[0].runs[0].font.name == "Amazon Ember"
    assert run_font.text_frame.paragraphs[0].runs[0].font.size.pt == 26


def test_header_image_replaced_in_place(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    header = visible(prs)[2]
    pics = [sh for sh in header.shapes if sh.shape_type == 13]
    assert len(pics) == 1
    assert pics[0].width == prs.slide_width and pics[0].height == prs.slide_height
    # same layer: still the last shape, over the placeholders
    assert list(header.shapes)[-1].shape_id == pics[0].shape_id
    assert pics[0].image.size == (160, 90)
    assert "hosted by RBI" in texts(header)


def test_without_header_image_keeps_the_old_one(setup, tmp_path):
    prs, report = run(setup, tmp_path, header=False)
    pics = [sh for sh in visible(prs)[2].shapes if sh.shape_type == 13]
    assert len(pics) == 1
    assert any("header slide" in t for t in report.todo)


def test_shoutout_names_until_a_logo_exists(setup, tmp_path):
    prs, report = run(setup, tmp_path)
    shout = visible(prs)[1]
    assert not [sh for sh in shout.shapes if sh.shape_type == 13]
    assert "Raiffeisen Bank International AG" in texts(shout)
    assert any("#Todo" in t and "logo" in t for t in texts(shout))
    assert any("sponsor logo missing" in t for t in report.todo)


def test_shoutout_with_logo(setup, tmp_path):
    template, event_file, stills, header = setup
    event = json.loads(event_file.read_text())
    (event_file.parent / "assets").mkdir()
    Image.new("RGB", (400, 100), "red").save(event_file.parent / "assets" / "rbi.png")
    event["sponsors"][0]["logo"] = f"events/{event['id']}/assets/rbi.png"
    event_file.write_text(json.dumps(event))
    out = tmp_path / "out.pptx"
    report = build(Inputs(template, event_file, stills, header), out)
    shout = visible(Presentation(str(out)))[1]
    pics = [sh for sh in shout.shapes if sh.shape_type == 13]
    assert len(pics) == 1 and pics[0].image.size == (400, 100)
    assert abs(pics[0].width / pics[0].height - 4) < 0.01
    assert not any("#Todo" in t for t in texts(shout))
    assert not report.todo


def test_see_you_soon_points_at_the_next_meetup(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    see = visible(prs)[-1]
    assert "at our next Meetup\x0bawscommunity.at" in texts(see)


def test_hidden_slides_untouched(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    hidden = [s for s in prs.slides if s._element.get("show") == "0"]
    assert len(hidden) == 1
    assert texts(hidden[0]) == ["Meetup #47"]


def test_talks_are_pictures_the_rest_native_in_the_deck_font(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    v = visible(prs)
    for slide in (v[5], v[8]):  # the two talks (the host's slot is between talk one and the break)
        shapes = list(slide.shapes)
        assert len(shapes) == 1 and shapes[0].shape_type == 13
        assert (shapes[0].left, shapes[0].top, shapes[0].width, shapes[0].height) == (0, 0, prs.slide_width, prs.slide_height)
    for slide, first in ((v[3], "Agenda"), (v[7], "Break"), (v[9], "Thank you")):
        assert texts(slide)[0] == first
        fonts = {r.font.name for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
        assert fonts == {"Amazon Ember"}
    assert "Back at 19:20" in texts(v[7])
    assert any("Talk two" in t for t in texts(v[7]))
    assert any("Networking until 22:00" in t for t in texts(v[9]))


def test_agenda_lists_every_segment(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    agenda = " | ".join(texts(visible(prs)[3]))
    for when in ("17:00-18:00", "18:10-18:50", "18:50-19:20", "20:00-22:00"):
        assert when in agenda
    assert "Roman Jokel, Mladen Trampic" in agenda


def test_next_community_day_slide(setup, tmp_path):
    template, event_file, stills, header = setup
    event = json.loads(event_file.read_text())
    people = event_file.parent.parent.parent / "people"
    people.mkdir()
    Image.new("RGB", (300, 300), "orange").save(people / "linda-mohamed.png")
    (event_file.parent.parent.parent / "community").mkdir()
    Image.new("RGB", (1280, 720), "teal").save(event_file.parent.parent.parent / "community" / "communityday-adria.jpg")
    event["upcoming"] = [
        {"name": "AWS Community Day Adria", "date": "2026-10-22", "location": "Ljubljana, Slovenia", "url": "https://awscommunityadria.com/",
         "speakers": [{"name": "Linda Mohamed", "photo": "people/linda-mohamed.png"}, {"name": "Darko Mesaros"}],
         "screenshot": "community/communityday-adria.jpg"},
        {"name": "AWS Community Day Switzerland", "date": "2026-10-29", "location": "Zurich", "url": "https://www.awsug.ch/"},
    ]
    event_file.write_text(json.dumps(event))
    out = tmp_path / "out.pptx"
    build(Inputs(template, event_file, stills, header), out)
    v = visible(Presentation(str(out)))
    slide = v[10]
    t = " | ".join(texts(slide))
    assert texts(slide)[0] == "Next AWS Community Day"
    assert "AWS Community Day Adria" in t and "Thursday, 22 October 2026  |  Ljubljana, Slovenia" in t
    assert "Linda Mohamed" in t and "Darko Mesaros" in t
    assert "Then AWS Community Day Switzerland, 29 October" in t
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    # Linda's round photo (Darko has none here), the website screenshot and the QR code
    assert len(pics) == 3
    assert [p.click_action.hyperlink.address for p in pics if p.click_action.hyperlink.address] == ["https://awscommunityadria.com/"] * 2
    url_box = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("awscommunityadria.com"))
    assert url_box.click_action.hyperlink.address == "https://awscommunityadria.com/"
    assert texts(v[11])[0] == "See you soon"


def test_missing_still_is_an_error(setup, tmp_path):
    template, event_file, stills, header = setup
    (stills / "SpeakerHandover--talk-2.png").unlink()
    with pytest.raises(FileNotFoundError, match="make stills"):
        build(Inputs(template, event_file, stills, header), tmp_path / "out.pptx")
    assert not (tmp_path / "out.pptx").exists()


def test_cli(setup, tmp_path, capsys):
    template, event_file, stills, header = setup
    out = tmp_path / "deck.pptx"
    code = main(["--template", str(template), "--event", str(event_file), "--stills", str(stills), "--header", str(header), "--out", str(out)])
    assert code == 0 and out.is_file()
    printed = capsys.readouterr().out
    assert "visible slides in running order" in printed and "Meetup #53" in printed
    assert main(["--template", str(tmp_path / "nope.pptx"), "--event", str(event_file), "--stills", str(stills), "--out", str(out)]) == 2


def test_evening_slides_follow_the_agenda(setup):
    _, event_file, stills, _ = setup
    ev = evening_slides(json.loads(event_file.read_text()), stills)
    assert [e.kind for e in ev] == ["image", "break", "image", "closing"]
    assert [e.image.stem for e in ev if e.image] == ["SpeakerHandover--talk-1", "SpeakerHandover--talk-2"]


def test_svg_logo_uses_the_png_beside_it(setup, tmp_path):
    template, event_file, stills, header = setup
    event = json.loads(event_file.read_text())
    assets = event_file.parent / "assets"
    assets.mkdir()
    (assets / "rbi.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>")
    event["sponsors"][0]["logo"] = "assets/rbi.svg"
    event_file.write_text(json.dumps(event))
    out = tmp_path / "out.pptx"
    report = build(Inputs(template, event_file, stills, header), out)
    assert any("is an SVG; add rbi.png" in t for t in report.todo)
    Image.new("RGB", (300, 100), "red").save(assets / "rbi.png")
    report = build(Inputs(template, event_file, stills, header), out)
    shout = visible(Presentation(str(out)))[1]
    assert [sh.image.size for sh in shout.shapes if sh.shape_type == 13] == [(300, 100)]
    assert not report.todo


def test_canva_images_take_the_talk_and_header_slides(setup, tmp_path):
    template, event_file, stills, header = setup
    canva = tmp_path / "canva"
    canva.mkdir()
    Image.new("RGB", (192, 108), "green").save(canva / "talk-1.png")
    Image.new("RGB", (192, 108), "yellow").save(canva / "header.png")
    out = tmp_path / "out.pptx"
    build(Inputs(template, event_file, stills, header, canva), out)
    prs = Presentation(str(out))
    pic = lambda s: next(sh for sh in s.shapes if sh.shape_type == 13)
    v = visible(prs)
    assert pic(v[2]).image.blob == (canva / "header.png").read_bytes()
    assert pic(v[5]).image.blob == (canva / "talk-1.png").read_bytes()
    # talk 2 has no Canva image: the stream still
    assert pic(v[8]).image.blob == (stills / "SpeakerHandover--talk-2.png").read_bytes()


def test_whole_deck_in_amazon_ember(setup, tmp_path):
    prs, _ = run(setup, tmp_path)
    fonts = {r.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.font.name}
    assert fonts == {"Amazon Ember"}


def test_keep_fonts(setup, tmp_path):
    template, event_file, stills, header = setup
    out = tmp_path / "out.pptx"
    build(Inputs(template, event_file, stills, header, None, None), out)
    welcome = visible(Presentation(str(out)))[0]
    assert {r.font.name for sh in welcome.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs} == {"PT Sans"}


def test_what_we_do_links_each_group_and_adds_qr_codes(setup, tmp_path):
    from pptx.util import Inches
    from conftest import png, text
    template, event_file, stills, header = setup
    prs = Presentation(str(template))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, "What we do", 0.3, 1.4, 4.4)
    for name, x, y in (("vienna", 5.2, 0.24), ("women", 7.8, 0.24), ("linz", 6.26, 1.54)):
        s.shapes.add_picture(str(png(tmp_path / f"{name}.png", "white")), Inches(x), Inches(y), Inches(2.4), Inches(1.0))
    prs.save(str(template))
    event = json.loads(event_file.read_text())
    groups = [
        {"name": "AWS User Group Vienna", "url": "https://www.meetup.com/vienna/"},
        {"name": "AWS Women's User Group Vienna", "url": "https://www.meetup.com/women/"},
        {"name": "AWS User Group Linz", "url": "https://www.meetup.com/linz/"},
    ]
    event["community"] = {"name": "AWS User Group Vienna", "team": [], "groups": groups, "badges": [], "photos": [], "links": {}}
    event_file.write_text(json.dumps(event))
    out = tmp_path / "out.pptx"
    build(Inputs(template, event_file, stills, header), out)
    slide = next(sl for sl in visible(Presentation(str(out))) if "What we do" in texts(sl))
    links = [p.click_action.hyperlink.address for p in slide.shapes if p.shape_type == 13]
    # the three group pages, top left, top right, middle; then the three QR codes
    assert links[:3] == [g["url"] for g in groups]
    assert links[3:] == [g["url"] for g in groups]
    assert [t for t in texts(slide) if t.endswith(("Vienna", "Linz")) and t != "What we do"] == ["UG Vienna", "Women's UG Vienna", "UG Linz"]


def test_who_we_are_takes_the_organizer_photos_from_people(setup, tmp_path):
    from pptx.util import Inches
    from conftest import png, text
    template, event_file, stills, header = setup
    prs = Presentation(str(template))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, "Who we are", 0.3, 1.4, 3.5)
    old = s.shapes.add_picture(str(png(tmp_path / "old-linda.png", "gray", (100, 100))), Inches(4.7), Inches(0.75), Inches(0.93), Inches(0.93))
    old._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom").set("prst", "ellipse")
    text(s, "Linda Mohamed\nlinda@example.org", 4.3, 1.87, 1.6, 0.85, 10)
    prs.save(str(template))
    people = event_file.parent.parent.parent / "people"
    people.mkdir(exist_ok=True)
    Image.new("RGB", (800, 800), "green").save(people / "linda-mohamed.png")
    event = json.loads(event_file.read_text())
    event["community"] = {"name": "AWS User Group Vienna", "team": [{"name": "Linda Mohamed", "email": "linda@awscommunity.at"}], "groups": [], "badges": [], "photos": [], "links": {}}
    event_file.write_text(json.dumps(event))
    out = tmp_path / "out.pptx"
    report = build(Inputs(template, event_file, stills, header), out)
    slide = next(sl for sl in visible(Presentation(str(out))) if "Who we are" in texts(sl))
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    assert pic.image.size == (800, 800)
    # still the round picture, now centred in the right part of the slide (one person: in its middle)
    assert pic._element.xpath(".//a:prstGeom/@prst") == ["ellipse"]
    right_mid = Inches(4) + Inches(6) / 2
    assert abs((pic.left + pic.width / 2) - right_mid) < Inches(0.02)
    assert "Linda Mohamed\nlinda@awscommunity.at" in texts(slide)
    assert not any("linda@example.org" in t for t in texts(slide))
    assert any("team laid out evenly" in c for c in report.changed)
